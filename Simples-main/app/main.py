from io import BytesIO
import os
from pathlib import Path
from typing import Any
from fastapi import Depends, FastAPI, File, HTTPException, Response, UploadFile
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import load_workbook
from pydantic import BaseModel
from sqlalchemy.orm import Session
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage
from .database import Base, SessionLocal, engine, get_db
from .models import Aluno, Resultado, Turma, Usuario
from .scanner import ScanError, scan_card
from .auth import admin_user, can_access_turma, clear_login_cookie, current_user, hash_password, set_login_cookie, verify_password

BASE_DIR = Path(__file__).resolve().parent
Base.metadata.create_all(bind=engine)

def bootstrap_admin():
    login = os.getenv('ADMIN_LOGIN', '').strip().lower()
    password = os.getenv('ADMIN_PASSWORD', '')
    name = os.getenv('ADMIN_NAME', 'Administrador').strip() or 'Administrador'
    if not login or not password: return
    db = SessionLocal()
    try:
        if not db.query(Usuario).filter(Usuario.login == login).first():
            db.add(Usuario(nome=name, login=login, senha_hash=hash_password(password), tipo='admin', ativo=True)); db.commit()
    finally: db.close()
bootstrap_admin()

app = FastAPI(title='ScoreView')
app.mount('/static', StaticFiles(directory=BASE_DIR / 'static'), name='static')
@app.get('/')
def home(): return FileResponse(BASE_DIR / 'static' / 'index.html')
@app.get('/health')
def health(): return {'ok': True}

class LoginIn(BaseModel): login: str; senha: str
class UserIn(BaseModel): nome: str; login: str; senha: str; tipo: str = 'professor'; turma_ids: list[int] = []
class UserUpdate(BaseModel): nome: str; login: str; senha: str | None = None; tipo: str = 'professor'; ativo: bool = True; turma_ids: list[int] = []

@app.post('/api/auth/login')
def login(payload: LoginIn, response: Response, db: Session = Depends(get_db)):
    user = db.query(Usuario).filter(Usuario.login == payload.login.strip().lower()).first()
    if not user or not user.ativo or not verify_password(payload.senha, user.senha_hash): raise HTTPException(401, 'Usuário ou senha inválidos.')
    set_login_cookie(response, user.id); return {'ok': True, 'nome': user.nome, 'tipo': user.tipo}
@app.post('/api/auth/logout')
def logout(response: Response): clear_login_cookie(response); return {'ok': True}
@app.get('/api/auth/me')
def me(user: Usuario = Depends(current_user)): return {'id': user.id, 'nome': user.nome, 'login': user.login, 'tipo': user.tipo}

@app.post('/api/scan')
async def scan(file: UploadFile = File(...), user: Usuario = Depends(current_user)):
    try: return scan_card(await file.read())
    except ScanError as exc: return JSONResponse(status_code=422, content={'detail': str(exc)})
    except Exception: return JSONResponse(status_code=500, content={'detail': 'Não foi possível processar a foto. Tente outra imagem.'})

@app.get('/api/turmas')
def listar_turmas(db: Session = Depends(get_db), user: Usuario = Depends(current_user)):
    turmas = db.query(Turma).order_by(Turma.nome).all() if user.tipo == 'admin' else sorted(user.turmas, key=lambda t:t.nome)
    return [{'id': t.id, 'nome': t.nome, 'total_alunos': len(t.alunos)} for t in turmas]
@app.get('/api/turmas/{turma_id}/alunos')
def listar_alunos(turma_id: int, db: Session = Depends(get_db), user: Usuario = Depends(current_user)):
    turma = db.get(Turma, turma_id)
    if not turma: raise HTTPException(404, 'Turma não encontrada.')
    if not can_access_turma(user, turma_id): raise HTTPException(403, 'Você não tem acesso a esta turma.')
    alunos = db.query(Aluno).filter(Aluno.turma_id == turma_id).order_by(Aluno.nome).all()
    return [{'id': a.id, 'nome': a.nome, 'matricula': a.matricula} for a in alunos]

@app.post('/api/alunos/importar')
async def importar_alunos(file: UploadFile = File(...), db: Session = Depends(get_db), user: Usuario = Depends(admin_user)):
    if not (file.filename or '').lower().endswith(('.xlsx', '.xlsm')): raise HTTPException(400, 'Envie uma planilha .xlsx.')
    try:
        wb=load_workbook(BytesIO(await file.read()),read_only=True,data_only=True); rows=list(wb.active.iter_rows(values_only=True))
    except Exception: raise HTTPException(400, 'Não foi possível ler a planilha.')
    if not rows: raise HTTPException(400, 'A planilha está vazia.')
    def norm(v): return str(v or '').strip().lower().replace('í','i').replace('á','a').replace('ã','a').replace('ç','c')
    headers=[norm(v) for v in rows[0]]; aliases={'nome':['nome','nome do aluno','aluno'],'turma':['turma','sala'],'matricula':['matricula','matrícula','registro','ra']}; indexes={}
    for field,names in aliases.items():
        for i,h in enumerate(headers):
            if h in [norm(n) for n in names]: indexes[field]=i; break
    if len(indexes)!=3: raise HTTPException(400,'A primeira linha precisa conter as colunas Nome do aluno, Turma e Matrícula.')
    criados=atualizados=ignorados=0
    for row in rows[1:]:
        def value(field):
            i=indexes[field]; return str(row[i] if i<len(row) and row[i] is not None else '').strip()
        nome,turma_nome,matricula=value('nome'),value('turma'),value('matricula')
        if not nome or not turma_nome or not matricula: ignorados+=1; continue
        turma=db.query(Turma).filter(Turma.nome==turma_nome).first()
        if not turma: turma=Turma(nome=turma_nome); db.add(turma); db.flush()
        aluno=db.query(Aluno).filter(Aluno.matricula==matricula).first()
        if aluno: aluno.nome=nome; aluno.turma_id=turma.id; atualizados+=1
        else: db.add(Aluno(nome=nome,matricula=matricula,turma_id=turma.id)); criados+=1
    db.commit(); return {'criados':criados,'atualizados':atualizados,'ignorados':ignorados}

@app.get('/api/admin/usuarios')
def list_users(db: Session=Depends(get_db), _:Usuario=Depends(admin_user)):
    users=db.query(Usuario).order_by(Usuario.nome).all(); return [{'id':u.id,'nome':u.nome,'login':u.login,'tipo':u.tipo,'ativo':u.ativo,'turma_ids':[t.id for t in u.turmas]} for u in users]
@app.post('/api/admin/usuarios')
def create_user(payload:UserIn, db:Session=Depends(get_db), _:Usuario=Depends(admin_user)):
    login=payload.login.strip().lower()
    if len(payload.senha)<6: raise HTTPException(400,'A senha deve ter pelo menos 6 caracteres.')
    if payload.tipo not in ('admin','professor'): raise HTTPException(400,'Tipo inválido.')
    if db.query(Usuario).filter(Usuario.login==login).first(): raise HTTPException(409,'Este login já está cadastrado.')
    turmas=db.query(Turma).filter(Turma.id.in_(payload.turma_ids)).all() if payload.turma_ids else []
    u=Usuario(nome=payload.nome.strip(),login=login,senha_hash=hash_password(payload.senha),tipo=payload.tipo,ativo=True,turmas=turmas); db.add(u); db.commit(); db.refresh(u); return {'id':u.id}
@app.put('/api/admin/usuarios/{user_id}')
def update_user(user_id:int,payload:UserUpdate,db:Session=Depends(get_db),admin:Usuario=Depends(admin_user)):
    u=db.get(Usuario,user_id)
    if not u: raise HTTPException(404,'Usuário não encontrado.')
    login=payload.login.strip().lower(); existing=db.query(Usuario).filter(Usuario.login==login,Usuario.id!=user_id).first()
    if existing: raise HTTPException(409,'Este login já está cadastrado.')
    if payload.tipo not in ('admin','professor'): raise HTTPException(400,'Tipo inválido.')
    if user_id==admin.id and (payload.tipo!='admin' or not payload.ativo): raise HTTPException(400,'Você não pode remover seu próprio acesso de administrador.')
    u.nome=payload.nome.strip();u.login=login;u.tipo=payload.tipo;u.ativo=payload.ativo
    if payload.senha:
        if len(payload.senha)<6: raise HTTPException(400,'A senha deve ter pelo menos 6 caracteres.')
        u.senha_hash=hash_password(payload.senha)
    u.turmas=db.query(Turma).filter(Turma.id.in_(payload.turma_ids)).all() if payload.turma_ids else []
    db.commit(); return {'ok':True}



def _pdf_resultado(r: Resultado, turma: Turma) -> BytesIO:
    out = BytesIO()
    doc = SimpleDocTemplate(out, pagesize=A4, rightMargin=16*mm, leftMargin=16*mm, topMargin=14*mm, bottomMargin=14*mm)
    styles = getSampleStyleSheet()
    title = ParagraphStyle('SchoolTitle', parent=styles['Title'], alignment=TA_CENTER, fontSize=15, leading=18, spaceAfter=3)
    sub = ParagraphStyle('SchoolSub', parent=styles['Normal'], alignment=TA_CENTER, fontSize=9, textColor=colors.HexColor('#555555'))
    story = []
    logo = BASE_DIR / 'static' / 'logo-escola.png'
    if logo.exists():
        story += [RLImage(str(logo), width=35*mm, height=35*mm), Spacer(1, 2*mm)]
    story += [Paragraph('Colégio Estadual em Período Integral João Barbosa Reis', title), Paragraph('CEPI-JBR - Relatório de Resultado', sub), Spacer(1, 7*mm)]
    info = [
        ['Aluno', r.aluno.nome], ['Matrícula', r.aluno.matricula], ['Turma', turma.nome],
        ['Prova / bloco', r.prova_nome], ['Questões', str(r.quantidade_questoes)],
        ['Acertos', str(r.acertos)], ['Erros', str(r.erros)], ['Anuladas', str(r.anuladas)],
        ['Em branco', str(r.em_branco)], ['Nota', f'{r.nota:.2f}'.replace('.', ',')],
        ['Data', r.criado_em.strftime('%d/%m/%Y %H:%M')],
    ]
    table = Table(info, colWidths=[42*mm, 120*mm])
    table.setStyle(TableStyle([
        ('BACKGROUND',(0,0),(0,-1),colors.HexColor('#f1f3f5')), ('FONTNAME',(0,0),(0,-1),'Helvetica-Bold'),
        ('GRID',(0,0),(-1,-1),0.5,colors.HexColor('#d6d9dc')), ('VALIGN',(0,0),(-1,-1),'MIDDLE'),
        ('LEFTPADDING',(0,0),(-1,-1),7), ('RIGHTPADDING',(0,0),(-1,-1),7), ('TOPPADDING',(0,0),(-1,-1),7), ('BOTTOMPADDING',(0,0),(-1,-1),7),
    ]))
    story += [table, Spacer(1, 8*mm), Paragraph('Documento gerado pelo ScoreView. Relatório para conferência e uso administrativo.', sub)]
    doc.build(story)
    out.seek(0)
    return out

@app.delete('/api/admin/usuarios/{user_id}')
def delete_user(user_id:int, db:Session=Depends(get_db), admin:Usuario=Depends(admin_user)):
    u=db.get(Usuario,user_id)
    if not u: raise HTTPException(404,'Usuário não encontrado.')
    if u.id==admin.id: raise HTTPException(400,'Você não pode excluir o usuário administrador que está em uso.')
    if u.tipo!='professor': raise HTTPException(400,'Por segurança, somente cadastros de professor podem ser excluídos por esta opção.')
    db.delete(u); db.commit(); return {'ok':True}

@app.delete('/api/resultados/{resultado_id}')
def delete_resultado(resultado_id:int, db:Session=Depends(get_db), user:Usuario=Depends(current_user)):
    r=db.get(Resultado,resultado_id)
    if not r: raise HTTPException(404,'Resultado não encontrado.')
    if not can_access_turma(user,r.aluno.turma_id): raise HTTPException(403,'Você não tem acesso a este resultado.')
    db.delete(r); db.commit(); return {'ok':True}

@app.get('/api/resultados/{resultado_id}/pdf')
def resultado_pdf(resultado_id:int, db:Session=Depends(get_db), user:Usuario=Depends(current_user)):
    r=db.get(Resultado,resultado_id)
    if not r: raise HTTPException(404,'Resultado não encontrado.')
    if not can_access_turma(user,r.aluno.turma_id): raise HTTPException(403,'Você não tem acesso a este resultado.')
    pdf=_pdf_resultado(r,r.aluno.turma)
    safe=''.join(c if c.isalnum() or c in '-_' else '_' for c in f'{r.aluno.nome}_{r.prova_nome}')[:100]
    return StreamingResponse(pdf, media_type='application/pdf', headers={'Content-Disposition':f'attachment; filename="relatorio_{safe}.pdf"'})

class ResultadoIn(BaseModel):
    aluno_id:int; prova_nome:str='Prova'; quantidade_questoes:int; gabarito:list[str]; questions:list[dict[str,Any]]
@app.post('/api/resultados')
def salvar_resultado(payload:ResultadoIn,db:Session=Depends(get_db),user:Usuario=Depends(current_user)):
    aluno=db.get(Aluno,payload.aluno_id)
    if not aluno: raise HTTPException(404,'Aluno não encontrado.')
    if not can_access_turma(user,aluno.turma_id): raise HTTPException(403,'Você não tem acesso a este aluno.')
    count=payload.quantidade_questoes
    if count not in (20,30,45) or len(payload.gabarito)!=count: raise HTTPException(400,'Quantidade/gabarito inválido.')
    qs=sorted(payload.questions,key=lambda q:int(q.get('number',0)))[:count]
    if len(qs)<count: raise HTTPException(400,'Leitura incompleta do cartão.')
    correct=mult=blank=0;answers=[]
    for i,q in enumerate(qs):
        status=q.get('status');answer=q.get('answer') if status=='OK' else None;answers.append(answer or ('MULT' if status=='MULT' else ''))
        if status=='MULT':mult+=1
        elif status=='BLANK':blank+=1
        if status=='OK' and answer==payload.gabarito[i]:correct+=1
    wrong=count-correct;nota=round(correct/count*10,2)
    result=Resultado(aluno_id=aluno.id,prova_nome=payload.prova_nome.strip() or 'Prova',quantidade_questoes=count,gabarito=','.join(payload.gabarito),respostas=','.join(answers),acertos=correct,erros=wrong,anuladas=mult,em_branco=blank,nota=nota)
    db.add(result);db.commit();db.refresh(result);return {'id':result.id,'acertos':correct,'erros':wrong,'anuladas':mult,'em_branco':blank,'nota':nota}

def _turma_resultados(db: Session, turma_id: int, prova_nome: str | None = None):
    query = db.query(Resultado).join(Aluno).filter(Aluno.turma_id == turma_id)
    if prova_nome:
        query = query.filter(Resultado.prova_nome == prova_nome)
    return query.order_by(Aluno.nome, Resultado.criado_em.desc()).all()

def _diagnostico_turma(resultados):
    if not resultados:
        return {'total_resultados':0,'alunos_com_resultado':0,'media_nota':0,'aproveitamento':0,'media_acertos':0,'media_erros':0,'media_brancos':0,'media_anuladas':0,'distribuicao':[],'questoes':[],'texto':'Ainda não há resultados suficientes para gerar um diagnóstico.'}
    alunos=len({r.aluno_id for r in resultados})
    media_nota=sum(r.nota for r in resultados)/len(resultados)
    media_acertos=sum(r.acertos for r in resultados)/len(resultados)
    media_erros=sum(r.erros for r in resultados)/len(resultados)
    media_brancos=sum(r.em_branco for r in resultados)/len(resultados)
    media_anuladas=sum(r.anuladas for r in resultados)/len(resultados)
    aproveitamento=sum((r.acertos/r.quantidade_questoes*100) if r.quantidade_questoes else 0 for r in resultados)/len(resultados)
    faixas={'0–39%':0,'40–59%':0,'60–79%':0,'80–100%':0}
    for r in resultados:
        pct=(r.acertos/r.quantidade_questoes*100) if r.quantidade_questoes else 0
        if pct<40: faixas['0–39%']+=1
        elif pct<60: faixas['40–59%']+=1
        elif pct<80: faixas['60–79%']+=1
        else: faixas['80–100%']+=1
    questoes=[]
    if len({r.prova_nome for r in resultados})==1:
        maxq=max((r.quantidade_questoes for r in resultados),default=0)
        for i in range(maxq):
            attempted=correct=blank=multi=0
            for r in resultados:
                if i>=r.quantidade_questoes: continue
                rs=(r.respostas or '').split(','); gs=(r.gabarito or '').split(',')
                ans=rs[i].strip().upper() if i<len(rs) else ''; key=gs[i].strip().upper() if i<len(gs) else ''
                attempted+=1
                if ans=='MULT': multi+=1
                elif not ans: blank+=1
                elif key and ans==key: correct+=1
            if attempted:
                questoes.append({'questao':i+1,'respostas':attempted,'acertos':correct,'percentual':round(correct/attempted*100,1),'brancos':blank,'anuladas':multi})
    fortes=sorted(questoes,key=lambda x:(-x['percentual'],x['questao']))[:3]
    atencao=sorted(questoes,key=lambda x:(x['percentual'],x['questao']))[:3]
    if questoes:
        fortes_txt=', '.join(f"Q{x['questao']} ({x['percentual']:.0f}%)" for x in fortes)
        atencao_txt=', '.join(f"Q{x['questao']} ({x['percentual']:.0f}%)" for x in atencao)
        texto=(f"A turma apresentou aproveitamento médio de {aproveitamento:.1f}%. Como pontos de maior domínio, destacam-se {fortes_txt}. "
               f"As questões que merecem retomada ou reforço são {atencao_txt}. A média de questões em branco foi {media_brancos:.1f} por resultado e a média de anuladas foi {media_anuladas:.1f}.")
    else:
        texto=(f"A turma apresentou aproveitamento médio de {aproveitamento:.1f}%. A média foi {media_nota:.2f} e houve, em média, {media_brancos:.1f} questões em branco por resultado. "
               "Para um diagnóstico por questão, selecione uma única prova no relatório.")
    return {'total_resultados':len(resultados),'alunos_com_resultado':alunos,'media_nota':round(media_nota,2),'aproveitamento':round(aproveitamento,1),'media_acertos':round(media_acertos,2),'media_erros':round(media_erros,2),'media_brancos':round(media_brancos,2),'media_anuladas':round(media_anuladas,2),'distribuicao':[{'faixa':k,'quantidade':v} for k,v in faixas.items()],'questoes':questoes,'texto':texto}

def _pdf_relatorio_turma(turma, resultados, diagnostico, prova_nome=None):
    out=BytesIO(); doc=SimpleDocTemplate(out,pagesize=A4,rightMargin=12*mm,leftMargin=12*mm,topMargin=12*mm,bottomMargin=12*mm)
    styles=getSampleStyleSheet(); title=ParagraphStyle('RTTitle',parent=styles['Title'],alignment=TA_CENTER,fontSize=16,leading=19,spaceAfter=3); sub=ParagraphStyle('RTSub',parent=styles['Normal'],alignment=TA_CENTER,fontSize=9,textColor=colors.HexColor('#555555')); h=ParagraphStyle('RTH',parent=styles['Heading2'],fontSize=12,leading=15,spaceBefore=7,spaceAfter=5); small=ParagraphStyle('RTSmall',parent=styles['Normal'],fontSize=8.5,leading=12)
    story=[]; logo=BASE_DIR/'static'/'logo-escola.png'
    if logo.exists(): story += [RLImage(str(logo),width=25*mm,height=25*mm),Spacer(1,1*mm)]
    story += [Paragraph('Colégio Estadual em Período Integral João Barbosa Reis',title),Paragraph('CEPI-JBR - Relatório e diagnóstico de desempenho da turma',sub),Spacer(1,5*mm)]
    info=[['Turma',turma.nome],['Prova',prova_nome or 'Todas as provas'],['Resultados',str(diagnostico['total_resultados'])],['Alunos com resultado',str(diagnostico['alunos_com_resultado'])],['Média da turma',f"{diagnostico['media_nota']:.2f}".replace('.',',')],['Aproveitamento médio',f"{diagnostico['aproveitamento']:.1f}%"]]
    t=Table(info,colWidths=[45*mm,120*mm]); t.setStyle(TableStyle([('BACKGROUND',(0,0),(0,-1),colors.HexColor('#f1f3f5')),('FONTNAME',(0,0),(0,-1),'Helvetica-Bold'),('GRID',(0,0),(-1,-1),.4,colors.HexColor('#d6d9dc')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('PADDING',(0,0),(-1,-1),6)])); story += [t,Spacer(1,6*mm),Paragraph('Diagnóstico geral',h),Paragraph(diagnostico['texto'],small),Spacer(1,4*mm)]
    dist=[['Faixa de aproveitamento','Resultados']]+[[x['faixa'],str(x['quantidade'])] for x in diagnostico['distribuicao']]
    td=Table(dist,colWidths=[95*mm,40*mm]); td.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e9ecef')),('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'),('GRID',(0,0),(-1,-1),.4,colors.HexColor('#d6d9dc')),('ALIGN',(1,1),(-1,-1),'CENTER'),('PADDING',(0,0),(-1,-1),5)])); story += [Paragraph('Distribuição dos resultados',h),td]
    if diagnostico['questoes']:
        story += [Paragraph('Desempenho por questão',h)]; qt=[['Questão','Respostas','Acertos','Aproveitamento','Brancos','Anuladas']]+[[str(x['questao']),str(x['respostas']),str(x['acertos']),f"{x['percentual']:.1f}%",str(x['brancos']),str(x['anuladas'])] for x in diagnostico['questoes']]; qtb=Table(qt,colWidths=[18*mm,24*mm,22*mm,32*mm,24*mm,24*mm],repeatRows=1); qtb.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e9ecef')),('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'),('GRID',(0,0),(-1,-1),.35,colors.HexColor('#d6d9dc')),('ALIGN',(0,0),(-1,-1),'CENTER'),('FONTSIZE',(0,0),(-1,-1),7.5),('PADDING',(0,0),(-1,-1),4)])); story += [qtb]
    story += [Spacer(1,7*mm),Paragraph('Resultados individuais',h)]; rows=[['Aluno','Matrícula','Prova','Acertos','Erros','Nota']]+[[r.aluno.nome,r.aluno.matricula,r.prova_nome,f'{r.acertos}/{r.quantidade_questoes}',str(r.erros),f'{r.nota:.2f}'.replace('.',',')] for r in resultados]
    if len(rows)==1: rows.append(['Nenhum resultado','','','','',''])
    rt=Table(rows,colWidths=[50*mm,28*mm,37*mm,24*mm,18*mm,18*mm],repeatRows=1); rt.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e9ecef')),('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'),('GRID',(0,0),(-1,-1),.35,colors.HexColor('#d6d9dc')),('FONTSIZE',(0,0),(-1,-1),7.5),('PADDING',(0,0),(-1,-1),4)])); story += [rt,Spacer(1,5*mm),Paragraph('Diagnóstico gerado automaticamente a partir dos resultados salvos. Use-o como apoio ao planejamento do feedback e das próximas atividades.',small)]
    doc.build(story); out.seek(0); return out

@app.get('/api/turmas/{turma_id}/provas')
def listar_provas_turma(turma_id:int, db:Session=Depends(get_db), user:Usuario=Depends(current_user)):
    turma=db.get(Turma,turma_id)
    if not turma: raise HTTPException(404,'Turma não encontrada.')
    if not can_access_turma(user,turma_id): raise HTTPException(403,'Você não tem acesso a esta turma.')
    return [x[0] for x in db.query(Resultado.prova_nome).join(Aluno).filter(Aluno.turma_id==turma_id).distinct().order_by(Resultado.prova_nome).all()]

@app.get('/api/turmas/{turma_id}/relatorio')
def relatorio(turma_id:int,prova_nome:str|None=None,db:Session=Depends(get_db),user:Usuario=Depends(current_user)):
    turma=db.get(Turma,turma_id)
    if not turma: raise HTTPException(404,'Turma não encontrada.')
    if not can_access_turma(user,turma_id): raise HTTPException(403,'Você não tem acesso a esta turma.')
    resultados=_turma_resultados(db,turma_id,prova_nome); diagnostico=_diagnostico_turma(resultados)
    return {'turma':turma.nome,'prova':prova_nome or 'Todas as provas','diagnostico':diagnostico,'resultados':[{'id':r.id,'aluno':r.aluno.nome,'matricula':r.aluno.matricula,'prova':r.prova_nome,'questoes':r.quantidade_questoes,'acertos':r.acertos,'erros':r.erros,'anuladas':r.anuladas,'em_branco':r.em_branco,'nota':r.nota,'data':r.criado_em.isoformat()} for r in resultados]}

@app.get('/api/turmas/{turma_id}/relatorio/pdf')
def relatorio_turma_pdf(turma_id:int,prova_nome:str|None=None,db:Session=Depends(get_db),user:Usuario=Depends(current_user)):
    turma=db.get(Turma,turma_id)
    if not turma: raise HTTPException(404,'Turma não encontrada.')
    if not can_access_turma(user,turma_id): raise HTTPException(403,'Você não tem acesso a esta turma.')
    resultados=_turma_resultados(db,turma_id,prova_nome); diagnostico=_diagnostico_turma(resultados); pdf=_pdf_relatorio_turma(turma,resultados,diagnostico,prova_nome)
    safe=''.join(c if c.isalnum() or c in '-_' else '_' for c in f'{turma.nome}_{prova_nome or "todas"}')[:100]
    return StreamingResponse(pdf,media_type='application/pdf',headers={'Content-Disposition':f'inline; filename="relatorio_turma_{safe}.pdf"'})
